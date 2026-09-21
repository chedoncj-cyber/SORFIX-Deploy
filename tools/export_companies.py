"""
Export all listed companies across GSE, JSE, NGX, NSE to Excel.
Run: python tools/export_companies.py
"""
import platform
import subprocess
import sys
import json
import urllib.request
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
except ImportError:
    print("pip install openpyxl"); sys.exit(1)

from tools.companies_data import COMPANIES

BASE_URL = "http://localhost:9090"

VENUE_META = {
    "GSE": {"name": "Ghana Stock Exchange",        "currency": "GHS", "fee_bps": 15,  "latency_ms": 2,  "fx_bps": 12, "dark": "1B5E20", "light": "E8F5E9", "mid": "A5D6A7"},
    "JSE": {"name": "Johannesburg Stock Exchange", "currency": "ZAR", "fee_bps": 20,  "latency_ms": 8,  "fx_bps":  8, "dark": "0D47A1", "light": "E3F2FD", "mid": "90CAF9"},
    "NGX": {"name": "Nigerian Exchange Group",     "currency": "NGN", "fee_bps": 30,  "latency_ms": 12, "fx_bps": 20, "dark": "B71C1C", "light": "FFEBEE", "mid": "EF9A9A"},
    "NSE": {"name": "Nairobi Securities Exchange", "currency": "KES", "fee_bps": 21,  "latency_ms": 15, "fx_bps": 15, "dark": "4A148C", "light": "F3E5F5", "mid": "CE93D8"},
}


def api(path):
    try:
        with urllib.request.urlopen(f"{BASE_URL}{path}", timeout=4) as r:
            return json.loads(r.read())
    except Exception:
        return None


def border(color="CCCCCC"):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def hfont(size=9, bold=True, color="FFFFFF"):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def cfont(size=9, bold=False, color="212121"):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def center():
    return Alignment(horizontal="center", vertical="center")


def left(indent=1):
    return Alignment(horizontal="left", vertical="center", indent=indent)


def right():
    return Alignment(horizontal="right", vertical="center")


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ─── fetch live books ───────────────────────────────────────────────────────────
def fetch_live_data():
    live = {}
    syms_resp = api("/market-data/symbols")
    if not syms_resp:
        return live
    for venue, syms in syms_resp.items():
        for sym in syms:
            book = api(f"/market-data/book/{venue}/{sym}")
            if book:
                bid_qty = sum(lvl["quantity"] for lvl in book.get("bids", []))
                ask_qty = sum(lvl["quantity"] for lvl in book.get("asks", []))
                live[f"{venue}:{sym}"] = {
                    "best_bid": book.get("best_bid"),
                    "best_ask": book.get("best_ask"),
                    "spread_bps": round((book.get("spread") or 0) * 10_000, 1),
                    "mid": round(((book.get("best_bid") or 0) + (book.get("best_ask") or 0)) / 2, 4),
                    "bid_depth": bid_qty,
                    "ask_depth": ask_qty,
                }
    return live


# ─── Title row helper ────────────────────────────────────────────────────────────
def title_row(ws, text, cols, dark_color, row=1, height=28):
    ws.merge_cells(f"A{row}:{get_column_letter(cols)}{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
    c.fill = fill(dark_color)
    c.alignment = left(2)
    ws.row_dimensions[row].height = height


def header_row(ws, headers, dark_color, row=2, height=20):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = hfont(9)
        c.fill = fill(dark_color)
        c.alignment = center()
        c.border = border()
    ws.row_dimensions[row].height = height


# ─── Build workbook ─────────────────────────────────────────────────────────────
def build_workbook(live: dict) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    now_str = datetime.now().strftime("%d %b %Y  %H:%M")
    total   = sum(len(v) for v in COMPANIES.values())

    # ── Sheet 1: All Companies ────────────────────────────────────────────────
    ws = wb.active
    ws.title = "All Companies"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    title_row(ws,
        f"Pan-African Listed Companies — GSE · JSE · NGX · NSE   ({total} companies)   |   {now_str}",
        16, "1A237E")

    hdrs = ["#", "Exchange", "Symbol", "Company Name", "Sector",
            "Currency", "Ref. Price", "Best Bid", "Best Ask", "Mid Price",
            "Spread (bps)", "Bid Depth", "Ask Depth",
            "Taker Fee (bps)", "Latency (ms)", "Live"]
    cws  = [5, 12, 14, 32, 16, 10, 12, 12, 12, 12, 13, 12, 12, 15, 13, 8]
    header_row(ws, hdrs, "283593")
    set_col_widths(ws, cws)

    row = 3
    counter = 1
    for venue, companies in COMPANIES.items():
        m = VENUE_META[venue]
        for i, (sym, d) in enumerate(companies.items()):
            key  = f"{venue}:{sym}"
            lv   = live.get(key, {})
            is_live = bool(lv)
            bg = m["light"] if i % 2 == 0 else "FAFAFA"

            vals = [
                counter, venue, sym, d["name"], d["sector"],
                m["currency"], d["price"],
                lv.get("best_bid"), lv.get("best_ask"), lv.get("mid"),
                lv.get("spread_bps"), lv.get("bid_depth"), lv.get("ask_depth"),
                m["fee_bps"], m["latency_ms"],
                "● LIVE" if is_live else "○",
            ]

            for col, val in enumerate(vals, 1):
                c = ws.cell(row=row, column=col, value=val)
                c.border = border()
                c.font   = cfont(9)
                c.fill   = fill(bg)

                if col == 1:                          # #
                    c.alignment = center()
                    c.font = cfont(9, color="999999")
                elif col == 2:                        # exchange badge
                    c.fill = fill(m["dark"])
                    c.font = hfont(9)
                    c.alignment = center()
                elif col in (3,):                     # symbol
                    c.font = cfont(9, bold=True, color="1A237E")
                    c.alignment = left(1)
                elif col == 4:                        # company name
                    c.alignment = left(1)
                elif col == 5:                        # sector
                    c.alignment = left(1)
                    c.font = cfont(8, color="555555")
                elif col in (7, 8, 9, 10):            # prices
                    c.alignment = right()
                    c.number_format = '#,##0.0000'
                elif col == 11:                       # spread bps
                    c.alignment = right()
                    c.number_format = '0.0'
                elif col in (12, 13):                 # depth
                    c.alignment = right()
                    c.number_format = '#,##0'
                elif col in (14, 15):                 # fee/latency
                    c.alignment = center()
                elif col == 16:                       # live indicator
                    c.alignment = center()
                    c.font = cfont(9, bold=is_live,
                                   color="22C55E" if is_live else "AAAAAA")
                else:
                    c.alignment = center()

            ws.row_dimensions[row].height = 15
            row += 1
            counter += 1

        # Thin divider between exchanges
        for col in range(1, 17):
            c = ws.cell(row=row, column=col, value="")
            c.fill = fill("E0E0E0")
        ws.row_dimensions[row].height = 3
        row += 1

    # Colour-scale on Spread (bps) col K
    ws.conditional_formatting.add(
        f"K3:K{row}", ColorScaleRule(
            start_type="min",  start_color="63BE7B",
            mid_type="num",    mid_value=20, mid_color="FFEB84",
            end_type="max",    end_color="F8696B"))

    # ── Sheets 2-5: Per-exchange ──────────────────────────────────────────────
    for venue, companies in COMPANIES.items():
        m   = VENUE_META[venue]
        ws2 = wb.create_sheet(f"{venue} — {m['currency']}")
        ws2.sheet_view.showGridLines = False
        ws2.freeze_panes = "A3"

        title_row(ws2,
            f"{m['name']}  ({venue})  ·  {m['currency']}  ·  {len(companies)} companies  ·  Fee {m['fee_bps']}bps  ·  Latency {m['latency_ms']}ms",
            14, m["dark"])

        hdrs2 = ["#", "Symbol", "Company Name", "Sector",
                 "Ref. Price", "Best Bid", "Best Ask", "Mid Price",
                 "Spread (bps)", "Bid Depth", "Ask Depth", "Bid/Ask Ratio",
                 "Taker Fee (bps)", "Live"]
        cws2  = [5, 14, 32, 16, 12, 12, 12, 12, 13, 12, 12, 13, 15, 8]
        header_row(ws2, hdrs2, m["dark"])
        set_col_widths(ws2, cws2)

        for i, (sym, d) in enumerate(companies.items()):
            key  = f"{venue}:{sym}"
            lv   = live.get(key, {})
            is_live = bool(lv)
            bg = m["light"] if i % 2 == 0 else "FAFAFA"
            ba_ratio = round(lv["bid_depth"] / lv["ask_depth"], 3) if lv.get("ask_depth") else None

            vals2 = [
                i + 1, sym, d["name"], d["sector"],
                d["price"],
                lv.get("best_bid"), lv.get("best_ask"), lv.get("mid"),
                lv.get("spread_bps"), lv.get("bid_depth"), lv.get("ask_depth"),
                ba_ratio, m["fee_bps"],
                "● LIVE" if is_live else "○",
            ]

            for col, val in enumerate(vals2, 1):
                c = ws2.cell(row=i + 3, column=col, value=val)
                c.border = border()
                c.font   = cfont(9)
                c.fill   = fill(bg)

                if col == 1:
                    c.alignment = center(); c.font = cfont(9, color="999999")
                elif col == 2:
                    c.font = cfont(9, bold=True, color=m["dark"]); c.alignment = left(1)
                elif col in (3,):
                    c.alignment = left(1)
                elif col == 4:
                    c.alignment = left(1); c.font = cfont(8, color="555555")
                elif col in (5, 6, 7, 8):
                    c.alignment = right(); c.number_format = '#,##0.0000'
                elif col == 9:
                    c.alignment = right(); c.number_format = '0.0'
                elif col in (10, 11):
                    c.alignment = right(); c.number_format = '#,##0'
                elif col == 12:
                    c.alignment = center(); c.number_format = '0.000'
                elif col == 13:
                    c.alignment = center()
                elif col == 14:
                    c.alignment = center()
                    c.font = cfont(9, bold=is_live, color="22C55E" if is_live else "AAAAAA")
                else:
                    c.alignment = center()
            ws2.row_dimensions[i + 3].height = 15

        # Colour scale on spread col I
        ws2.conditional_formatting.add(
            f"I3:I{len(companies)+3}", ColorScaleRule(
                start_type="min",  start_color="63BE7B",
                mid_type="num",    mid_value=20, mid_color="FFEB84",
                end_type="max",    end_color="F8696B"))

    # ── Sheet 6: Sector Summary ───────────────────────────────────────────────
    ws6 = wb.create_sheet("Sector Breakdown")
    ws6.sheet_view.showGridLines = False
    title_row(ws6, f"Sector Breakdown by Exchange   |   {now_str}", 6, "37474F")
    header_row(ws6, ["Exchange", "Sector", "Count", "% of Exchange", "Ref. Avg Price", "Currency"], "455A64")
    set_col_widths(ws6, [14, 22, 9, 15, 18, 12])

    row6 = 3
    for venue, companies in COMPANIES.items():
        m = VENUE_META[venue]
        sectors: dict = {}
        for d in companies.values():
            s = d["sector"]
            if s not in sectors:
                sectors[s] = {"count": 0, "prices": []}
            sectors[s]["count"] += 1
            sectors[s]["prices"].append(d["price"])

        total_v = len(companies)
        for j, (sec, sv) in enumerate(sorted(sectors.items(), key=lambda x: -x[1]["count"])):
            bg = m["light"] if j % 2 == 0 else "FAFAFA"
            avg_price = round(sum(sv["prices"]) / len(sv["prices"]), 4)
            vals6 = [venue, sec, sv["count"], round(sv["count"]/total_v*100, 1), avg_price, m["currency"]]
            for col, val in enumerate(vals6, 1):
                c = ws6.cell(row=row6, column=col, value=val)
                c.border = border()
                c.fill   = fill(bg)
                c.font   = cfont(9)
                if col == 1:
                    c.fill = fill(m["dark"]); c.font = hfont(9); c.alignment = center()
                elif col == 4:
                    c.alignment = center(); c.number_format = '0.0'
                elif col == 5:
                    c.alignment = right(); c.number_format = '#,##0.0000'
                elif col == 3:
                    c.alignment = center()
                else:
                    c.alignment = left(1)
            ws6.row_dimensions[row6].height = 15
            row6 += 1

        for col in range(1, 7):
            c = ws6.cell(row=row6, column=col, value="")
            c.fill = fill("E0E0E0")
        ws6.row_dimensions[row6].height = 3
        row6 += 1

    # ── Sheet 7: Exchange Summary ─────────────────────────────────────────────
    ws7 = wb.create_sheet("Exchange Summary")
    ws7.sheet_view.showGridLines = False
    title_row(ws7, "Exchange Overview", 8, "1A237E")
    header_row(ws7,
        ["Exchange", "Full Name", "Currency", "Listed Companies",
         "Taker Fee (bps)", "Latency (ms)", "FX Cost (bps)", "FX Rank"],
        "283593")
    set_col_widths(ws7, [14, 32, 10, 18, 16, 14, 14, 10])

    for rank, (venue, m) in enumerate(sorted(VENUE_META.items(), key=lambda x: x[1]["fx_bps"]), 1):
        c_count = len(COMPANIES[venue])
        row7 = rank + 2
        vals7 = [venue, m["name"], m["currency"], c_count,
                 m["fee_bps"], m["latency_ms"], m["fx_bps"], rank]
        for col, val in enumerate(vals7, 1):
            c = ws7.cell(row=row7, column=col, value=val)
            c.border = border()
            c.fill   = fill(m["light"])
            c.font   = cfont(9)
            if col == 1:
                c.fill = fill(m["dark"]); c.font = hfont(9); c.alignment = center()
            elif col == 2:
                c.alignment = left(1)
            else:
                c.alignment = center()
        ws7.row_dimensions[row7].height = 20

    return wb


def main():
    print(f"Fetching live market data from {BASE_URL}...")
    live = fetch_live_data()
    live_count = len(live)
    total = sum(len(v) for v in COMPANIES.values())
    print(f"  {live_count} live order books  /  {total} total companies")

    wb = build_workbook(live)

    # Resolve output directory: prefer OneDrive Desktop, fall back to plain Desktop,
    # then fall back to current working directory (works on any OS).
    _candidates = [
        Path.home() / "OneDrive" / "Desktop",
        Path.home() / "Desktop",
        Path.cwd(),
    ]
    out_dir = next((p for p in _candidates if p.exists()), Path.cwd())
    out_path = out_dir / f"FIX_SOR_AllCompanies_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    wb.save(str(out_path))
    print(f"\nSaved: {out_path}")
    # Open in file explorer — Windows only
    if platform.system() == "Windows":
        subprocess.Popen(["explorer", str(out_path)])
    return str(out_path)


if __name__ == "__main__":
    main()
