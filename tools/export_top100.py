"""
PAN SORFIX — Export Pan-African Listed Top 100 Companies to Excel.
Generates a formatted .xlsx on the user's Desktop.
"""

import sys
import os
from pathlib import Path
from datetime import date
from collections import Counter

sys.path.insert(0, str(Path(__file__).parent.parent))

from tools.companies_data import COMPANIES
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Exchange metadata ──────────────────────────────────────────────────────────
EXCHANGE_META = {
    "GSE":   {"flag": "🇬🇭", "full": "Ghana Stock Exchange",              "country": "Ghana",             "currency": "GHS", "color": "006B3F", "light": "E6F4EE", "settlement": "T+2"},
    "JSE":   {"flag": "🇿🇦", "full": "Johannesburg Stock Exchange",       "country": "South Africa",      "currency": "ZAR", "color": "007A4D", "light": "E6F4EE", "settlement": "T+3"},
    "NGX":   {"flag": "🇳🇬", "full": "Nigerian Exchange Group",           "country": "Nigeria",           "currency": "NGN", "color": "008751", "light": "E6F5ED", "settlement": "T+2"},
    "NSE":   {"flag": "🇰🇪", "full": "Nairobi Securities Exchange",       "country": "Kenya",             "currency": "KES", "color": "006600", "light": "ECF7EC", "settlement": "T+3"},
    "CSE":   {"flag": "🇲🇦", "full": "Casablanca Stock Exchange",         "country": "Morocco",           "currency": "MAD", "color": "C1272D", "light": "FDEEEE", "settlement": "T+2"},
    "EGX":   {"flag": "🇪🇬", "full": "Egyptian Exchange",                 "country": "Egypt",             "currency": "EGP", "color": "CE1126", "light": "FDECEA", "settlement": "T+2"},
    "BRVM":  {"flag": "🌍",   "full": "BRVM (Bourse Régionale UEMOA)",    "country": "West Africa (8)",   "currency": "XOF", "color": "FF6600", "light": "FFF3E0", "settlement": "T+3"},
    "BSE":   {"flag": "🇧🇼", "full": "Botswana Stock Exchange",           "country": "Botswana",          "currency": "BWP", "color": "75AADB", "light": "EAF3FC", "settlement": "T+3"},
    "NSX":   {"flag": "🇳🇦", "full": "Namibia Stock Exchange",            "country": "Namibia",           "currency": "NAD", "color": "003580", "light": "E8EEF8", "settlement": "T+3"},
    "SEM":   {"flag": "🇲🇺", "full": "Stock Exchange of Mauritius",       "country": "Mauritius",         "currency": "MUR", "color": "EA2839", "light": "FDECEA", "settlement": "T+3"},
    "MSE":   {"flag": "🇲🇼", "full": "Malawi Stock Exchange",             "country": "Malawi",            "currency": "MWK", "color": "000000", "light": "F2F2F2", "settlement": "T+3"},
    "BVMT":  {"flag": "🇹🇳", "full": "Bourse des Valeurs Mobilières de Tunis", "country": "Tunisia",    "currency": "TND", "color": "E70013", "light": "FDEEEE", "settlement": "T+2"},
    "DSE":   {"flag": "🇹🇿", "full": "Dar es Salaam Stock Exchange",      "country": "Tanzania",          "currency": "TZS", "color": "1EB53A", "light": "E9F9EC", "settlement": "T+3"},
    "ZSE":   {"flag": "🇿🇼", "full": "Zimbabwe Stock Exchange",           "country": "Zimbabwe",          "currency": "ZWL", "color": "006400", "light": "E6F4E6", "settlement": "T+3"},
    "LUSE":  {"flag": "🇿🇲", "full": "Lusaka Securities Exchange",        "country": "Zambia",            "currency": "ZMW", "color": "198A00", "light": "E8F7E6", "settlement": "T+3"},
}

# Row alternating shades per exchange
ROW_COLORS = {
    "GSE":   ("C8E6C9", "DCEDC8"), "JSE":  ("B2DFDB", "CCE8E4"),
    "NGX":   ("C8E6C9", "D7EED7"), "NSE":  ("D0EDD0", "E2F2E2"),
    "CSE":   ("FDDEDE", "FEE9E9"), "EGX":  ("FBDCDC", "FCE8E8"),
    "BRVM":  ("FFE0B2", "FFECCC"), "BSE":  ("DDEEFF", "EAF3FF"),
    "NSX":   ("D5E8F8", "E5F0FA"), "SEM":  ("FDDEDE", "FEE9E9"),
    "MSE":   ("E8E8E8", "F2F2F2"), "BVMT": ("FDDEDE", "FEE9E9"),
    "DSE":   ("D4EDDA", "E5F5EA"), "ZSE":  ("C8E6C9", "D7EDD7"),
    "LUSE":  ("CBECCB", "DBF0DB"),
}

# ── Company picks: 100 total across 15 exchanges ───────────────────────────────
PICKS = {
    "JSE":   9,   # Largest African exchange by market cap
    "NGX":   8,   # Largest in West Africa
    "NSE":   8,   # Largest in East Africa
    "EGX":   8,   # North Africa
    "CSE":   7,   # Morocco
    "BRVM":  7,   # West Africa (regional)
    "SEM":   7,   # Mauritius
    "BVMT":  7,   # Tunisia
    "GSE":   6,   # Ghana
    "DSE":   6,   # Tanzania
    "BSE":   6,   # Botswana
    "ZSE":   6,   # Zimbabwe
    "LUSE":  6,   # Zambia
    "NSX":   5,   # Namibia
    "MSE":   4,   # Malawi (smallest)
}
assert sum(PICKS.values()) == 100, f"Total={sum(PICKS.values())}"

# ── Build rows ─────────────────────────────────────────────────────────────────
rows = []
for exchange, n in PICKS.items():
    meta    = EXCHANGE_META[exchange]
    co_dict = COMPANIES.get(exchange, {})
    for symbol, info in list(co_dict.items())[:n]:
        rows.append({
            "exchange":  exchange,
            "flag":      meta["flag"],
            "country":   meta["country"],
            "currency":  meta["currency"],
            "symbol":    symbol,
            "name":      info["name"],
            "sector":    info["sector"],
            "price":     info["price"],
            "vol_pct":   round(info["vol"] * 100, 3),
        })

# ── Helpers ────────────────────────────────────────────────────────────────────
def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color="000000", size=11, name="Calibri"):
    return Font(bold=bold, color=color, size=size, name=name)

def _border():
    s = Side(style="thin", color="D0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)

def _center():
    return Alignment(horizontal="center", vertical="center")

def _left():
    return Alignment(horizontal="left", vertical="center")

# ── Create workbook ────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()

# ════════════════════════════════════════════════════════════════════════════════
# Sheet 1 — Master List
# ════════════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Pan-African Top 100"

# Title block
ws.merge_cells("A1:K1")
tc = ws["A1"]
tc.value     = "PAN SORFIX  —  Pan-African Listed Top 100 Companies"
tc.font      = Font(bold=True, color="FFFFFF", size=16, name="Calibri")
tc.fill      = _fill("002B18")
tc.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 34

ws.merge_cells("A2:K2")
sc = ws["A2"]
sc.value = (
    f"GSE · JSE · NGX · NSE · CSE · EGX · BRVM · BSE · NSX · SEM · MSE · BVMT · DSE · ZSE · LUSE  |  "
    f"Generated: {date.today().strftime('%d %B %Y')}  |  FIX 4.2 Pan-African Smart Order Router"
)
sc.font      = Font(italic=True, color="A0C8A0", size=10, name="Calibri")
sc.fill      = _fill("003D20")
sc.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20
ws.row_dimensions[3].height = 6

# Column headers
HEADERS = ["#", "Flag", "Exchange", "Country", "Symbol",
           "Company Name", "Sector", "Currency", "Price (Local)", "Daily Vol %", "SOR Status"]
HDR_ROW = 4

for ci, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=HDR_ROW, column=ci, value=h)
    cell.font      = Font(bold=True, color="FFFFFF", size=10, name="Calibri")
    cell.fill      = _fill("1A4731")
    cell.alignment = _center()
    cell.border    = _border()
ws.row_dimensions[HDR_ROW].height = 24

# Data rows
for i, row in enumerate(rows, 1):
    dr   = HDR_ROW + i
    alt  = i % 2
    c1, c2 = ROW_COLORS[row["exchange"]]
    bg   = c1 if alt == 1 else c2

    values = [
        i,
        row["flag"],
        row["exchange"],
        row["country"],
        row["symbol"],
        row["name"],
        row["sector"],
        row["currency"],
        row["price"],
        row["vol_pct"],
        "ACTIVE",
    ]

    for ci, val in enumerate(values, 1):
        cell = ws.cell(row=dr, column=ci, value=val)
        cell.fill   = _fill(bg)
        cell.border = _border()

        if ci == 1:
            cell.font      = Font(bold=True, color="444444", size=10, name="Calibri")
            cell.alignment = _center()
        elif ci == 2:
            cell.font      = Font(size=13, name="Segoe UI Emoji")
            cell.alignment = _center()
        elif ci == 3:
            ex_color = EXCHANGE_META[row["exchange"]]["color"]
            cell.font      = Font(bold=True, color=ex_color, size=10, name="Calibri")
            cell.alignment = _center()
        elif ci == 5:
            cell.font      = Font(bold=True, color="003D20", size=10, name="Courier New")
            cell.alignment = _center()
        elif ci == 9:
            cell.number_format = '#,##0.00'
            cell.font          = Font(bold=True, color="1B5E20", size=10, name="Calibri")
            cell.alignment     = Alignment(horizontal="right", vertical="center")
        elif ci == 10:
            cell.number_format = '0.000"%"'
            vol = row["vol_pct"]
            color = "C62828" if vol >= 0.5 else ("E65100" if vol >= 0.3 else "2E7D32")
            cell.font      = Font(bold=True, color=color, size=10, name="Calibri")
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif ci == 11:
            cell.font  = Font(bold=True, color="1B5E20", size=9, name="Calibri")
            cell.fill  = _fill("DFFFDF") if alt == 1 else _fill("C8F5C8")
            cell.alignment = _center()
        else:
            cell.font      = Font(color="1A1A1A", size=10, name="Calibri")
            cell.alignment = _left()

    ws.row_dimensions[dr].height = 18

# Column widths
col_widths = [5, 6, 9, 26, 14, 44, 22, 10, 14, 12, 12]
for ci, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(ci)].width = w

ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HDR_ROW}:K{HDR_ROW + len(rows)}"

# ════════════════════════════════════════════════════════════════════════════════
# Sheet 2 — Exchange Summary
# ════════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Exchange Summary")

ws2.merge_cells("A1:I1")
ws2["A1"].value     = "PAN SORFIX — Exchange Summary (15 African Exchanges)"
ws2["A1"].font      = Font(bold=True, color="FFFFFF", size=14, name="Calibri")
ws2["A1"].fill      = _fill("002B18")
ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 30

ws2.merge_cells("A2:I2")
ws2["A2"].value     = f"Pan-African FIX 4.2 Smart Order Router  ·  {date.today().strftime('%d %B %Y')}"
ws2["A2"].font      = Font(italic=True, color="A0C8A0", size=10)
ws2["A2"].fill      = _fill("003D20")
ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[2].height = 18

SUM_HDRS = ["Exchange", "Full Name", "Flag", "Country", "Currency", "# in Top 100", "Avg Price", "Avg Vol %", "Settlement"]
for ci, h in enumerate(SUM_HDRS, 1):
    c = ws2.cell(row=4, column=ci, value=h)
    c.font      = Font(bold=True, color="FFFFFF", size=10, name="Calibri")
    c.fill      = _fill("1A4731")
    c.alignment = _center()
    c.border    = _border()
ws2.row_dimensions[4].height = 22

for ri, (ex, n) in enumerate(PICKS.items(), 1):
    meta    = EXCHANGE_META[ex]
    ex_rows = [r for r in rows if r["exchange"] == ex]
    avg_p   = sum(r["price"] for r in ex_rows) / len(ex_rows) if ex_rows else 0
    avg_v   = sum(r["vol_pct"] for r in ex_rows) / len(ex_rows) if ex_rows else 0
    bg      = ROW_COLORS[ex][0] if ri % 2 == 1 else ROW_COLORS[ex][1]

    vals = [ex, meta["full"], meta["flag"], meta["country"], meta["currency"],
            n, round(avg_p, 2), round(avg_v, 4), meta["settlement"]]
    for ci, v in enumerate(vals, 1):
        c = ws2.cell(row=4 + ri, column=ci, value=v)
        c.fill   = _fill(bg)
        c.border = _border()
        if ci == 1:
            c.font = Font(bold=True, color=meta["color"], size=11, name="Calibri")
            c.alignment = _center()
        elif ci == 3:
            c.font = Font(size=13, name="Segoe UI Emoji")
            c.alignment = _center()
        elif ci == 7:
            c.number_format = '#,##0.00'
            c.font = Font(bold=True, color="1B5E20", size=10, name="Calibri")
            c.alignment = Alignment(horizontal="right", vertical="center")
        elif ci == 8:
            c.number_format = '0.000"%"'
            c.font = Font(bold=True, color="B71C1C", size=10, name="Calibri")
            c.alignment = Alignment(horizontal="right", vertical="center")
        else:
            c.font = Font(color="1A1A1A", size=10, name="Calibri")
            c.alignment = _center()
    ws2.row_dimensions[4 + ri].height = 20

# Totals row
tot_row = 4 + len(PICKS) + 1
ws2.merge_cells(f"A{tot_row}:E{tot_row}")
ws2.cell(row=tot_row, column=1, value="TOTAL  (all 15 African exchanges)").font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
ws2.cell(row=tot_row, column=1).fill      = _fill("002B18")
ws2.cell(row=tot_row, column=1).alignment = Alignment(horizontal="center", vertical="center")
tot_n = ws2.cell(row=tot_row, column=6, value=100)
tot_n.font = Font(bold=True, color="FFFFFF", size=12, name="Calibri")
tot_n.fill = _fill("002B18")
tot_n.alignment = _center()
for ci in [7, 8, 9]:
    ws2.cell(row=tot_row, column=ci).fill = _fill("002B18")
ws2.row_dimensions[tot_row].height = 22

for ci, w in enumerate([11, 42, 7, 24, 10, 14, 14, 12, 12], 1):
    ws2.column_dimensions[get_column_letter(ci)].width = w

# ════════════════════════════════════════════════════════════════════════════════
# Sheet 3 — Sector Breakdown
# ════════════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Sector Breakdown")

sector_counts = Counter(r["sector"] for r in rows)
sector_ex     = {}
for r in rows:
    sector_ex.setdefault(r["sector"], set()).add(r["exchange"])

ws3.merge_cells("A1:E1")
ws3["A1"].value     = "PAN SORFIX — Sector Breakdown (Top 100 Companies)"
ws3["A1"].font      = Font(bold=True, color="FFFFFF", size=13, name="Calibri")
ws3["A1"].fill      = _fill("002B18")
ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 28

for ci, h in enumerate(["Sector", "# Companies", "% of 100", "Exchanges Present", "Bar"], 1):
    c = ws3.cell(row=3, column=ci, value=h)
    c.font = Font(bold=True, color="FFFFFF", size=10, name="Calibri")
    c.fill = _fill("1A4731")
    c.alignment = _center()
    c.border = _border()
ws3.row_dimensions[3].height = 22

SECTOR_COLORS = {
    "Banking":      "E3F2FD", "Mining":        "FFF9C4",
    "Technology":   "F3E5F5", "Consumer":      "E8F5E9",
    "Insurance":    "FCE4EC", "Energy":        "FFF3E0",
    "Telecom":      "E1F5FE", "Healthcare":    "FDECEA",
    "Agriculture":  "F9FBE7", "Manufacturing": "EFEBE9",
    "Financial":    "FFF8E1", "Real Estate":   "EDE7F6",
    "Materials":    "E0F2F1", "Industrial":    "FFF3E0",
    "Food & Bev":   "F1F8E9",
}

for ri, (sector, cnt) in enumerate(sorted(sector_counts.items(), key=lambda x: -x[1]), 1):
    bg  = SECTOR_COLORS.get(sector, "F5F5F5")
    pct = round(cnt / 100 * 100, 1)
    ex_list = ", ".join(sorted(sector_ex[sector]))
    bar = "█" * cnt + "░" * (20 - min(cnt, 20))

    for ci, v in enumerate([sector, cnt, pct, ex_list, bar], 1):
        c = ws3.cell(row=3 + ri, column=ci, value=v)
        c.fill   = _fill(bg)
        c.border = _border()
        if ci == 1:
            c.font = Font(bold=True, color="1A1A1A", size=10, name="Calibri")
            c.alignment = _left()
        elif ci in (2, 3):
            c.font = Font(bold=True, color="1B5E20", size=10, name="Calibri")
            c.alignment = _center()
        elif ci == 5:
            c.font = Font(color="1C6EAA", size=9, name="Courier New")
            c.alignment = _left()
        else:
            c.font = Font(color="333333", size=9, name="Calibri")
            c.alignment = _left()
    ws3.row_dimensions[3 + ri].height = 18

for ci, w in enumerate([24, 14, 12, 52, 28], 1):
    ws3.column_dimensions[get_column_letter(ci)].width = w

# ════════════════════════════════════════════════════════════════════════════════
# Sheet 4 — Country Index
# ════════════════════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("Country Index")

ws4.merge_cells("A1:F1")
ws4["A1"].value     = "PAN SORFIX — Country & Exchange Index"
ws4["A1"].font      = Font(bold=True, color="FFFFFF", size=13, name="Calibri")
ws4["A1"].fill      = _fill("002B18")
ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws4.row_dimensions[1].height = 28

for ci, h in enumerate(["Flag", "Country", "Exchange", "Currency", "Settlement", "# Companies in Export"], 1):
    c = ws4.cell(row=3, column=ci, value=h)
    c.font = Font(bold=True, color="FFFFFF", size=10, name="Calibri")
    c.fill = _fill("1A4731")
    c.alignment = _center()
    c.border = _border()
ws4.row_dimensions[3].height = 22

for ri, (ex, n) in enumerate(PICKS.items(), 1):
    meta = EXCHANGE_META[ex]
    bg   = ROW_COLORS[ex][0] if ri % 2 == 1 else ROW_COLORS[ex][1]
    for ci, v in enumerate([meta["flag"], meta["country"], ex, meta["currency"], meta["settlement"], n], 1):
        c = ws4.cell(row=3 + ri, column=ci, value=v)
        c.fill   = _fill(bg)
        c.border = _border()
        if ci == 1:
            c.font = Font(size=13, name="Segoe UI Emoji")
            c.alignment = _center()
        elif ci == 3:
            c.font = Font(bold=True, color=meta["color"], size=10, name="Calibri")
            c.alignment = _center()
        elif ci == 6:
            c.font = Font(bold=True, color="1B5E20", size=10, name="Calibri")
            c.alignment = _center()
        else:
            c.font = Font(color="1A1A1A", size=10, name="Calibri")
            c.alignment = _center()
    ws4.row_dimensions[3 + ri].height = 20

for ci, w in enumerate([7, 24, 10, 10, 12, 22], 1):
    ws4.column_dimensions[get_column_letter(ci)].width = w

# ── Save ───────────────────────────────────────────────────────────────────────
try:
    import winreg
    key = winreg.OpenKey(winreg.HKEY_CURRENT_USER,
        r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders")
    desktop_raw = winreg.QueryValueEx(key, "Desktop")[0]
    desktop = Path(os.path.expandvars(desktop_raw))
except Exception:
    desktop = Path.home() / "OneDrive" / "Desktop"

out_path = desktop / f"PanSORFIX_African_Top100_Companies_{date.today().strftime('%Y%m%d')}.xlsx"
wb.save(out_path)
print(f"Saved : {out_path}")
print(f"Rows  : {len(rows)} companies across {len(PICKS)} exchanges")
for ex, n in PICKS.items():
    print(f"  {ex:8s} {n} companies")
